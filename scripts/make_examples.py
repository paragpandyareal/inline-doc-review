#!/usr/bin/env python3
# Copyright (c) 2026 Parag Pandya. MIT License, see LICENSE.
# Lazy Panda Panel: https://github.com/paragpandyareal/lazy-panda-panel
"""Writes the sample Excel and Word files for /panda examples.

    make_examples.py <folder>

The samples are generated rather than shipped, so the plugin itself holds
only text files. Runs under the lazy-panda-panel venv (/panda setup), which
has openpyxl and python-docx. Prints the files it wrote, one per line.
"""
import os
import random
import sys

VENV = os.path.join(os.path.expanduser("~"), ".cache", "lazy-panda-panel", "venv")
VENV_PY = os.path.join(VENV, "bin", "python")


def is_trusted(path):
    """Owned by this user (or root), and writable by no one else (the user's own group aside): safe to run from."""
    try:
        st = os.stat(path)
    except OSError:
        return False
    group_ok = not st.st_mode & 0o020 or st.st_gid in (0, os.getgid())
    return st.st_uid in (0, os.getuid()) and not st.st_mode & 0o002 and group_ok


def venv_is_trusted():
    return all(is_trusted(p) for p in (VENV, os.path.join(VENV, "bin"), os.path.realpath(VENV_PY)))


def reexec_in_venv():
    # Compare the environment, not the interpreter: a venv's python links to the system one.
    # Only into a venv no one else can write to: otherwise read with the system Python.
    if os.path.exists(VENV_PY) and os.path.realpath(sys.prefix) != os.path.realpath(VENV) and venv_is_trusted():
        os.execv(VENV_PY, [VENV_PY, *sys.argv])


def budget(path):
    import openpyxl
    from openpyxl.styles import Font, PatternFill

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Budget"
    ws.append(["Line item", "Jan", "Feb", "Mar", "Total"])
    for name, a, b, c in [("Facebook ads", 8000, 8000, 8000), ("Letterbox print", 9600, 4800, 0),
                          ("Info nights", 3200, 3200, 3200), ("Installer bonus", 0, 5000, 5000)]:
        ws.append([name, a, b, c])
    for r in range(2, 6):
        ws[f"E{r}"] = f"=SUM(B{r}:D{r})"
    ws.append(["Total", "=SUM(B2:B5)", "=SUM(C2:C5)", "=SUM(D2:D5)", "=SUM(E2:E5)"])
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDEEDD")
    signups = wb.create_sheet("Sign-ups")
    signups.append(["Suburb", "Target", "Signed"])
    for row in [("Penrith", 200, 64), ("Blacktown", 180, 51), ("Mount Druitt", 120, 22)]:
        signups.append(row)
    wb.save(path)


def big_sheet(path):
    import openpyxl
    from openpyxl.styles import Font

    random.seed(7)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sites"
    months = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
    ws.append(["Site", "Region", "Tariff"] + months + ["Total", "Average", "Change"])
    regions = ["Penrith", "Blacktown", "Parramatta", "Liverpool", "Campbelltown"]
    for i in range(2, 62):
        ws.append([f"NMI-{4100 + i}", random.choice(regions), random.choice(["Flat", "TOU", "Demand"])]
                  + [random.randint(600, 2400) for _ in months])
        ws[f"P{i}"] = f"=SUM(D{i}:O{i})"
        ws[f"Q{i}"] = f"=AVERAGE(D{i}:O{i})"
        ws[f"R{i}"] = f"=(O{i}-D{i})/D{i}"
        ws[f"Q{i}"].number_format = "#,##0"
        ws[f"R{i}"].number_format = "0.0%"
    ws.append(["Total", "", ""] + [f"=SUM({c}2:{c}61)" for c in "DEFGHIJKLMNOP"])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    wb.save(path)


def proposal(path):
    import docx

    d = docx.Document()
    d.add_heading("Q1 Retail Energy Proposal", 0)
    d.add_heading("Summary", 1)
    d.add_paragraph("We propose a three-month pilot of the Green Saver Plan for 500 households in Western Sydney, "
                    "funded by the state solar rebate.")
    d.add_heading("Pricing", 1)
    d.add_paragraph("The plan charges 29.4c per kWh with a 98c daily supply charge. Customers with solar receive "
                    "a 6c feed-in tariff.")
    d.add_paragraph("A 5% pay-on-time discount applies to usage charges only.")
    table = d.add_table(rows=3, cols=3)
    table.style = "Table Grid"
    for r, row in enumerate([["Item", "Rate", "Unit"], ["Usage", "29.4", "c/kWh"], ["Supply", "98", "c/day"]]):
        for c, value in enumerate(row):
            table.cell(r, c).text = value
    d.add_heading("Timeline", 1)
    for step in ["Recruitment in January", "Installs in February", "First bills in March"]:
        d.add_paragraph(step, style="List Bullet")
    d.save(path)


def main():
    reexec_in_venv()
    folder = sys.argv[1]
    os.makedirs(folder, exist_ok=True)
    for name, make in [("pilot-budget.xlsx", budget), ("site-consumption.xlsx", big_sheet), ("energy-proposal.docx", proposal)]:
        path = os.path.join(folder, name)
        make(path)
        print(path)


if __name__ == "__main__":
    main()
