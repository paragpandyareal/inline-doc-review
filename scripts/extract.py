#!/usr/bin/env python3
# Copyright (c) 2026 Parag Pandya. MIT License, see LICENSE.
# Lazy Panda Panel: https://github.com/paragpandyareal/lazy-panda-panel
"""Reads a document for the Lazy Panda Panel and prints it as JSON.

    extract.py docx|pdf|xlsx <path>   one document, read only
    extract.py examples <folder>       writes the samples (make_examples.py)
    extract.py setup                   makes the venv the pane reads with

With no arguments it reads them from standard input, one per line (how the
pane sends them).

The libraries live in a venv of their own (~/.cache/lazy-panda-panel/venv),
installed from requirements.txt with exact versions and sha256 hashes, so
nothing is installed into the system Python. Each run re-executes itself
under that venv when it exists. Nothing here touches the network except
`setup`, and nothing writes to the document.
"""
import datetime
import json
import math
import os
import re
import subprocess
import sys
import zipfile

HERE = os.path.dirname(os.path.abspath(__file__))
VENV = os.path.join(os.path.expanduser("~"), ".cache", "lazy-panda-panel", "venv")
IS_WINDOWS = os.name == "nt"
VENV_BIN = os.path.join(VENV, "Scripts" if IS_WINDOWS else "bin")
VENV_PY = os.path.join(VENV_BIN, "python.exe" if IS_WINDOWS else "python")
REQUIREMENTS = os.path.join(HERE, "requirements.txt")

MAX_ROWS = 4000
MAX_SHEET_ROWS = 500
MAX_SHEET_COLS = 40
MAX_FILE_BYTES = 50 * 1024 * 1024
MAX_UNZIPPED_BYTES = 300 * 1024 * 1024

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
CONTROL = re.compile(r"\x1b\[[0-9;?]*[ -/]*[@-~]|\x1b\][^\x07\x1b]*(?:\x07|\x1b\\)|[\x00-\x08\x0b\x0c\x0e-\x1f\x7f-\x9f]")


def clean(text):
    """Text safe to draw: tabs as spaces, no terminal escape codes or control characters."""
    return CONTROL.sub("", str(text).replace("\t", "  ").replace("\r", ""))


def is_trusted(path):
    """Owned by this user (or root), and writable by no one else (the user's own group aside): safe to run from.

    Windows has no owner or mode bits here; its user profile folder is private to the user by default."""
    try:
        st = os.stat(path)
    except OSError:
        return False
    if IS_WINDOWS:
        return True
    group_ok = not st.st_mode & 0o020 or st.st_gid in (0, os.getgid())
    return st.st_uid in (0, os.getuid()) and not st.st_mode & 0o002 and group_ok


def venv_is_trusted():
    return all(is_trusted(p) for p in (VENV, VENV_BIN, os.path.realpath(VENV_PY)))


def reexec_in_venv(args):
    # Compare the environment, not the interpreter: a venv's python links to the system one.
    # Only into a venv no one else can write to: otherwise read with the system Python.
    if os.path.exists(VENV_PY) and os.path.realpath(sys.prefix) != os.path.realpath(VENV) and venv_is_trusted():
        argv = [VENV_PY, os.path.abspath(__file__), *args]
        if IS_WINDOWS:
            # Windows has no exec: the parent would exit at once and its output be lost, so wait for the child.
            sys.stdout.flush()
            sys.exit(subprocess.run(argv).returncode)
        os.execv(VENV_PY, argv)


def setup():
    os.umask(0o022)
    if not os.path.exists(VENV_PY):
        subprocess.run([sys.executable, "-m", "venv", VENV], check=True)
    # Hash-checked, wheels only: every file pip installs must match a hash in requirements.txt.
    subprocess.run([VENV_PY, "-m", "pip", "install", "--quiet", "--require-hashes", "--only-binary=:all:",
                    "--no-deps", "-r", REQUIREMENTS], check=True)
    return {"ok": True, "venv": VENV, "requirements": REQUIREMENTS}


def missing(package):
    return {
        "kind": "error",
        "message": f"Reading this file needs the Python package '{package}'. Run /panda setup once to install it.",
    }


def check_size(path, is_zip):
    """Refuses files too large to read safely, including zip bombs (docx and xlsx are zip files)."""
    if os.path.getsize(path) > MAX_FILE_BYTES:
        raise ValueError(f"the file is over {MAX_FILE_BYTES // (1024 * 1024)} MB, too large to show here")
    if is_zip:
        with zipfile.ZipFile(path) as archive:
            if sum(info.file_size for info in archive.infolist()) > MAX_UNZIPPED_BYTES:
                raise ValueError("the file expands to more than 300 MB, too large to show here")


def cut_note(kind):
    return f"Only the first {MAX_ROWS} rows of this {kind} are shown."


# ── Word ──

def run_spans(paragraph_element):
    """A paragraph's text as spans: bold, italic, strike, links. Includes tracked insertions; skips deletions."""
    spans = []
    for run in paragraph_element.iter(W + "r"):
        if any(ancestor.tag == W + "del" for ancestor in run.iterancestors()):
            continue
        text = "".join(t.text or "" for t in run.iter(W + "t"))
        if not text:
            continue
        props = run.find(W + "rPr")
        span = {"t": clean(text)}

        def on(tag):
            element = props.find(W + tag) if props is not None else None
            return element is not None and element.get(W + "val") not in ("0", "false")

        if on("b"):
            span["b"] = 1
        if on("i"):
            span["i"] = 1
        if on("strike"):
            span["s"] = 1
        if any(ancestor.tag == W + "hyperlink" for ancestor in run.iterancestors()):
            span["l"] = 1
        spans.append(span)
    return spans


def list_kind(document, paragraph_element):
    """For a numbered or bulleted paragraph: (depth, is_numbered); otherwise None."""
    num_pr = paragraph_element.find(f"{W}pPr/{W}numPr")
    if num_pr is None:
        return None
    level_el = num_pr.find(W + "ilvl")
    num_el = num_pr.find(W + "numId")
    depth = int(level_el.get(W + "val", "0")) if level_el is not None else 0
    num_id = num_el.get(W + "val") if num_el is not None else None
    is_numbered = False
    try:
        numbering = document.part.numbering_part.element
        abstract_id = numbering.xpath(f'./w:num[@w:numId="{num_id}"]/w:abstractNumId/@w:val')
        if abstract_id:
            fmt = numbering.xpath(f'./w:abstractNum[@w:abstractNumId="{abstract_id[0]}"]/w:lvl[@w:ilvl="{depth}"]/w:numFmt/@w:val')
            is_numbered = bool(fmt) and fmt[0] not in ("bullet", "none")
    except (KeyError, NotImplementedError, AttributeError, ValueError):
        pass
    return depth, is_numbered


def body_blocks(element):
    """The body's paragraphs and tables in order, looking inside content controls (w:sdt)."""
    for child in element.iterchildren():
        if child.tag == W + "sdt":
            content = child.find(W + "sdtContent")
            if content is not None:
                yield from body_blocks(content)
        elif child.tag in (W + "p", W + "tbl"):
            yield child


def docx_rows(path):
    try:
        import docx
        from docx.table import Table
        from docx.text.paragraph import Paragraph
    except ImportError:
        return missing("python-docx")

    check_size(path, is_zip=True)
    document = docx.Document(path)
    rows = []
    heading = None
    paragraph_index = 0
    table_index = 0
    numbers = {}
    has_title = False
    is_cut = False

    def space():
        if rows and rows[-1].get("style") != "space":
            rows.append({"text": "", "anchor": rows[-1]["anchor"], "style": "space", "unit": rows[-1]["unit"]})

    for block in body_blocks(document.element.body):
        if len(rows) >= MAX_ROWS:
            is_cut = True
            break
        if block.tag == W + "p":
            paragraph_index += 1
            paragraph = Paragraph(block, document)
            spans = run_spans(block)
            text = "".join(span["t"] for span in spans).strip()
            if not text:
                if block.find(f".//{W}drawing") is not None or block.find(f".//{W}pict") is not None:
                    space()
                    rows.append({"text": "[image]", "anchor": f"paragraph {paragraph_index} (an image)", "unit": "paragraph",
                                 "style": "p", "spans": [{"t": "▣ image", "d": 1}]})
                continue
            style = (paragraph.style.name if paragraph.style is not None else "") or ""
            # A document with a Title uses it as the top level; its Heading 1s sit one below.
            has_title = has_title or style == "Title"
            match = re.search(r"(\d+)", style)
            level = 1 if style == "Title" else min(3, (int(match.group(1)) if match else 3) + (1 if has_title else 0)) if style.startswith("Heading") else 0
            under = f' (under "{heading}")' if heading and not level else ""
            row = {"text": text, "anchor": f"paragraph {paragraph_index}{under}", "unit": "paragraph", "spans": spans}
            kind = list_kind(document, block)
            if kind is None and "List" in style:
                depth_match = re.search(r"(\d+)$", style)
                kind = (max(0, int(depth_match.group(1)) - 1) if depth_match else 0, "Number" in style)
            if level:
                heading = text
                numbers = {}
                space()
                row["style"] = f"h{level}"
                rows.append(row)
                if level == 1:
                    rows.append({"text": "", "anchor": row["anchor"], "style": "rule", "unit": "paragraph"})
            elif kind is not None:
                depth, is_numbered = kind
                for deeper in [d for d in numbers if d > depth]:
                    del numbers[deeper]  # a new parent item restarts the levels under it
                if rows and rows[-1].get("style") not in ("li", "space"):
                    space()
                if is_numbered:
                    numbers[depth] = numbers.get(depth, 0) + 1
                row.update({"style": "li", "indent": depth, "marker": f"{numbers[depth]}." if is_numbered else ["•", "◦", "▪"][depth % 3]})
                rows.append(row)
            else:
                numbers = {}
                space()
                row["style"] = "quote" if "Quote" in style else "p"
                rows.append(row)
        else:
            table_index += 1
            numbers = {}
            space()
            for row_index, table_row in enumerate(Table(block, document).rows, start=1):
                cells, seen = [], set()
                for cell in table_row.cells:
                    if id(cell._tc) in seen:  # a merged cell repeats the same element
                        continue
                    seen.add(id(cell._tc))
                    cells.append(clean(" ".join(cell.text.split())))
                rows.append({
                    "text": " | ".join(cells),
                    "cells": cells,
                    "isHeader": row_index == 1,
                    "anchor": f"table {table_index}, row {row_index}",
                    "unit": "table row",
                })
            space()
    while rows and rows[-1].get("style") == "space":
        rows.pop()
    doc = {"kind": "lines", "rows": rows, "isFormatted": True}
    if is_cut:
        doc["note"] = cut_note("document")
    return doc


# ── PDF ──

def pdf_rows(path):
    try:
        from pypdf import PdfReader
    except ImportError:
        return missing("pypdf")

    check_size(path, is_zip=False)
    reader = PdfReader(path)
    if reader.is_encrypted:
        try:
            if not reader.decrypt(""):
                return {"kind": "error", "message": "This PDF is password-protected, so its text can't be shown."}
        except Exception:
            return {"kind": "error", "message": "This PDF is password-protected, so its text can't be shown."}
    rows = []
    is_cut = False
    for page_number, page in enumerate(reader.pages, start=1):
        if len(rows) >= MAX_ROWS:
            is_cut = True
            break
        rows.append({"text": f"Page {page_number}", "anchor": f"page {page_number}", "style": "h3", "unit": "page"})
        for line_number, line in enumerate((page.extract_text() or "").splitlines(), start=1):
            if line.strip():
                rows.append({"text": clean(line.rstrip()), "anchor": f"page {page_number}, line {line_number}", "unit": "line"})
    doc = {"kind": "lines", "rows": rows}
    if is_cut:
        doc["note"] = f"Only the first {page_number - 1} of {len(reader.pages)} pages are shown."
    elif rows and all(row["unit"] == "page" for row in rows):
        doc["note"] = "This PDF has no text layer (it may be scanned), so only page markers are shown."
    return doc


# ── Excel ──

def number_text(value, fmt):
    """A number as the sheet's format would show it, roughly: %, currency, decimals, thousands, (negatives)."""
    fmt = fmt or "General"
    section = fmt.split(";")[0]
    if "." in section:
        decimals = len(re.match(r"[0#]*", section.split(".", 1)[1]).group(0))
    elif section == "General":
        if isinstance(value, float) and not value.is_integer():
            return f"{value:,.10g}" if abs(value) >= 0.001 else f"{value:.4g}"
        decimals = 0
    else:
        decimals = 0
    if "%" in fmt:
        return f"{value * 100:,.{decimals}f}%"
    text = f"{abs(value):,.{decimals}f}"
    symbol = next((s for s in ("$", "€", "£", "¥") if s in fmt), "")
    text = symbol + text
    if value < 0:
        return f"({text})" if "(" in fmt.split(";")[1 if ";" in fmt else 0] else f"-{text}"
    return text


def cell_text(value, fmt):
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, (int, float)):
        return number_text(value, fmt) if math.isfinite(value) else "#NUM!"
    if isinstance(value, datetime.datetime):
        if value.date() == datetime.date(1899, 12, 30) or (value.year == 1900 and "h" in (fmt or "").lower() and "y" not in (fmt or "").lower()):
            return value.strftime("%H:%M")
        return value.strftime("%Y-%m-%d") if (value.hour, value.minute, value.second) == (0, 0, 0) else value.strftime("%Y-%m-%d %H:%M")
    if isinstance(value, datetime.date):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, datetime.time):
        return value.strftime("%H:%M")
    if isinstance(value, datetime.timedelta):
        minutes = int(value.total_seconds() // 60)
        return f"{minutes // 60}:{minutes % 60:02d}"
    return clean(value).replace("\n", " ⏎ ")


def formula_text(raw):
    """A cell's formula as text, or None. Array formulas are objects in openpyxl."""
    text = getattr(raw, "text", raw)
    return clean(text) if isinstance(text, str) and text.startswith("=") else None


def xlsx_sheets(path):
    try:
        import openpyxl
        from openpyxl.utils import get_column_letter
    except ImportError:
        return missing("openpyxl")
    from formulas import Calculator

    check_size(path, is_zip=True)
    formulas = openpyxl.load_workbook(path, data_only=False)
    values = openpyxl.load_workbook(path, data_only=True)

    def raw_cell(sheet, row, col):
        if sheet not in formulas.sheetnames or row < 1 or col < 1:
            return None
        ws = formulas[sheet]
        if not hasattr(ws, "cell") or row > ws.max_row or col > ws.max_column:
            return None
        cached = values[sheet].cell(row=row, column=col).value
        raw = ws.cell(row=row, column=col).value
        text = formula_text(raw)
        return cached if (text is None or cached is not None) else text

    calculator = Calculator(raw_cell)
    not_calculated = 0
    sheets = []
    for f_sheet in formulas.worksheets:  # worksheets only: chart sheets have no cells
        name = f_sheet.title
        v_sheet = values[name]
        width = min(f_sheet.max_column or 1, MAX_SHEET_COLS)
        height = min(f_sheet.max_row or 1, MAX_SHEET_ROWS)
        rows = []
        for r, (f_row, v_row) in enumerate(zip(
                f_sheet.iter_rows(min_row=1, max_row=height, max_col=width),
                v_sheet.iter_rows(min_row=1, max_row=height, max_col=width, values_only=True)), start=1):
            cells = []
            for c, (f_cell, value) in enumerate(zip(f_row, v_row), start=1):
                raw = f_cell.value
                fmt = f_cell.number_format or "General"
                formula = formula_text(raw)
                if formula and value is None:
                    value = calculator.evaluate(name, r, c)
                    if value is None:
                        not_calculated += 1
                elif not formula:
                    value = raw
                cell = {"v": cell_text(value, fmt)}
                if formula:
                    cell["f"] = formula
                if isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(value):
                    cell["x"] = value
                cells.append(cell)
            rows.append(cells)
        while rows and all(not cell["v"] and "f" not in cell for cell in rows[-1]):
            rows.pop()
        sheet = {
            "name": clean(name),
            "cols": [get_column_letter(i) for i in range(1, width + 1)],
            "rows": rows,
            "isCut": (f_sheet.max_row or 0) > MAX_SHEET_ROWS or (f_sheet.max_column or 0) > MAX_SHEET_COLS,
        }
        if f_sheet.sheet_state != "visible":
            sheet["isHidden"] = True
        sheets.append(sheet)
    doc = {"kind": "grid", "sheets": sheets}
    if not_calculated:
        doc["note"] = (f"{not_calculated} formula cell{'s' if not_calculated != 1 else ''} use functions this pane can't calculate; "
                       "they show as ƒ with no value. Open the file in Excel to see them.")
    return doc


def main():
    # UTF-8 both ways, whatever the system's code page: paths and cells can hold any character.
    for stream in (sys.stdin, sys.stdout, sys.stderr):
        stream.reconfigure(encoding="utf-8")
    # The task and its path come as arguments, or on standard input, one per line (how the pane sends them).
    args = sys.argv[1:] or [line.strip() for line in sys.stdin.read().strip().split("\n", 1)]
    if args[:1] == ["setup"]:
        print(json.dumps(setup()))
        return
    reexec_in_venv(args)
    sys.path.insert(0, HERE)
    if len(args) == 2 and args[0] == "examples":
        import make_examples
        sys.exit(make_examples.write(args[1]))
    if len(args) != 2 or args[0] not in ("docx", "pdf", "xlsx"):
        print(json.dumps({"kind": "error", "message": "Usage: extract.py docx|pdf|xlsx|examples <path> (or on standard input, one per line)"}))
        return
    kind, path = args
    try:
        doc = {"docx": docx_rows, "pdf": pdf_rows, "xlsx": xlsx_sheets}[kind](path)
    except Exception as error:  # a corrupt, locked or oversized file: say so in the pane
        doc = {"kind": "error", "message": f"Could not read the file: {clean(error)}"}
    print(json.dumps(doc, allow_nan=False))


if __name__ == "__main__":
    main()
